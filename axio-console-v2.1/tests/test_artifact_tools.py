import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from core.code_tools import build_default_registry
from core.code_tools.artifact_tools import create_docx, create_pdf, create_pptx, create_xlsx
from core.code_tools.verification import extract_artifact_requirements


class ArtifactCreationTests(unittest.TestCase):
    def test_creates_and_reopens_native_office_and_pdf_artifacts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "brief.docx"
            xlsx_path = root / "data.xlsx"
            pptx_path = root / "deck.pptx"
            pdf_path = root / "report.pdf"

            docx = create_docx(
                str(docx_path),
                "Executive Brief",
                summary="Verified summary",
                sections=[{"heading": "Finding", "paragraphs": ["Evidence"], "bullets": ["Action"]}],
                sources=["https://example.com/source"],
            )
            xlsx = create_xlsx(
                str(xlsx_path),
                [{"name": "Summary", "title": "Metrics", "headers": ["Item", "Value"], "rows": [["Accuracy", 0.9], ["Total", "=SUM(B4:B4)"]]}],
            )
            pptx = create_pptx(
                str(pptx_path),
                "AXIO",
                subtitle="Artifact test",
                slides=[{"title": "Result", "bullets": ["Created", "Validated"]}],
                sources=["https://example.com/source"],
            )
            pdf = create_pdf(
                str(pdf_path),
                "AXIO Report",
                summary="Verified summary",
                sections=[{"heading": "Result", "paragraphs": ["Created"], "bullets": ["Validated"]}],
            )

            self.assertTrue(docx.ok, docx.content)
            self.assertTrue(xlsx.ok, xlsx.content)
            self.assertTrue(pptx.ok, pptx.content)
            self.assertTrue(pdf.ok, pdf.content)
            self.assertIn("Executive Brief", "\n".join(p.text for p in Document(docx_path).paragraphs))
            workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Summary"])
            workbook.close()
            self.assertEqual(len(Presentation(pptx_path).slides), 3)
            self.assertGreaterEqual(len(PdfReader(str(pdf_path)).pages), 1)

    def test_requested_artifact_blocks_completion_until_native_tool_validates_it(self):
        registry = build_default_registry()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target = root / "summary.docx"
            registry.start_task("crea un resumen ejecutivo en Word", [root])

            self.assertFalse(registry.completion_status()[0])
            blocked = registry.execute("verification_gate", {})
            self.assertIn("artifact:docx", blocked)

            created = registry.execute(
                "create_docx",
                {"path": str(target), "title": "Summary", "summary": "Completed"},
            )
            passed = registry.execute("verification_gate", {"required_checks": ["artifact:docx"]})

        self.assertIn("structurally validated", created)
        self.assertTrue(passed.startswith("VERIFICATION PASSED:"), passed)
        self.assertTrue(registry.completion_status()[0])

    def test_text_artifact_evidence_and_binary_extension_protection(self):
        registry = build_default_registry()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            registry.start_task("crea un archivo .txt", [root])
            text_path = root / "notes.txt"
            fake_word = root / "fake.docx"
            text = registry.execute("write_file", {"path": str(text_path), "content": "AXIO"})
            rejected = registry.execute("write_file", {"path": str(fake_word), "content": "not a DOCX"})
            passed = registry.execute("verification_gate", {})

        self.assertTrue(text.startswith("OK:"))
        self.assertTrue(rejected.startswith("ERROR:"))
        self.assertFalse(fake_word.exists())
        self.assertTrue(passed.startswith("VERIFICATION PASSED:"))

    def test_artifact_requirement_detection_is_task_scoped_and_multilingual(self):
        checks = extract_artifact_requirements(
            "Crea Word, Excel, PowerPoint, PDF y un archivo de texto .txt"
        )
        self.assertEqual(
            checks,
            ("artifact:docx", "artifact:pdf", "artifact:pptx", "artifact:txt", "artifact:xlsx"),
        )


if __name__ == "__main__":
    unittest.main()
