"""Safe creation of common Windows productivity artifacts with structural validation."""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path
from typing import Callable

from .registry import CodeTool, ToolResult


def _target(path: str, suffix: str, overwrite: bool) -> Path:
    target = Path(path).expanduser()
    if not target.is_absolute():
        raise ValueError(f"path must be absolute: {path}")
    target = target.resolve()
    if target.suffix.lower() != suffix:
        raise ValueError(f"path must end with {suffix}: {target}")
    if target.exists() and not overwrite:
        raise FileExistsError(f"File already exists; pass overwrite=true only when replacement is intended: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def _atomic_save(target: Path, save: Callable[[Path], None]) -> None:
    handle, temp_name = tempfile.mkstemp(prefix=f".{target.stem}.", suffix=target.suffix, dir=target.parent)
    os.close(handle)
    temporary = Path(temp_name)
    try:
        save(temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise ValueError("artifact writer produced an empty file")
        with temporary.open("r+b") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, target)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _clean_text(value: object, limit: int = 20_000) -> str:
    return str(value or "").strip()[:limit]


def create_docx(
    path: str,
    title: str,
    summary: str = "",
    sections: list[dict] | None = None,
    sources: list[str] | None = None,
    overwrite: bool = False,
) -> ToolResult:
    target = _target(path, ".docx", overwrite)
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return ToolResult("ERROR: DOCX creation requires python-docx. Install project requirements.", ok=False)

    normalized_sections = list(sections or [])[:40]
    normalized_sources = [_clean_text(item, 2_000) for item in (sources or [])[:50] if _clean_text(item)]

    def save(destination: Path) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10.5)
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = heading.add_run(_clean_text(title, 500) or "Untitled document")
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(31, 78, 121)
        if summary_text := _clean_text(summary):
            paragraph = document.add_paragraph(summary_text)
            paragraph.style = document.styles["Normal"]
        for item in normalized_sections:
            if not isinstance(item, dict):
                continue
            if heading_text := _clean_text(item.get("heading"), 500):
                document.add_heading(heading_text, level=1)
            paragraphs = item.get("paragraphs") or []
            if isinstance(paragraphs, str):
                paragraphs = [paragraphs]
            for paragraph_text in list(paragraphs)[:50]:
                if cleaned := _clean_text(paragraph_text):
                    document.add_paragraph(cleaned)
            bullets = item.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            for bullet in list(bullets)[:50]:
                if cleaned := _clean_text(bullet, 5_000):
                    document.add_paragraph(cleaned, style="List Bullet")
        if normalized_sources:
            document.add_heading("Sources", level=1)
            for source in normalized_sources:
                document.add_paragraph(source, style="List Bullet")
        document.save(destination)

    try:
        _atomic_save(target, save)
        reopened = Document(target)
        paragraph_count = len(reopened.paragraphs)
        if paragraph_count < 1:
            raise ValueError("DOCX validation found no paragraphs")
    except Exception as exc:
        target.unlink(missing_ok=True)
        return ToolResult(f"ERROR creating DOCX: {exc}", ok=False)
    return ToolResult(
        f"OK: Created and structurally validated DOCX ({paragraph_count} paragraphs): {target}",
        verification_ids=("artifact:docx",),
    )


def create_xlsx(
    path: str,
    sheets: list[dict],
    overwrite: bool = False,
) -> ToolResult:
    target = _target(path, ".xlsx", overwrite)
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return ToolResult("ERROR: XLSX creation requires openpyxl. Install project requirements.", ok=False)
    normalized_sheets = [item for item in list(sheets or [])[:25] if isinstance(item, dict)]
    if not normalized_sheets:
        return ToolResult("ERROR: sheets must contain at least one worksheet definition", ok=False)

    def save(destination: Path) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        used_names: set[str] = set()
        for index, spec in enumerate(normalized_sheets, 1):
            raw_name = _clean_text(spec.get("name"), 31) or f"Sheet{index}"
            name = re.sub(r"[\\/*?:\[\]]", "_", raw_name)[:31] or f"Sheet{index}"
            base_name = name
            counter = 2
            while name.casefold() in used_names:
                suffix = f"_{counter}"
                name = base_name[: 31 - len(suffix)] + suffix
                counter += 1
            used_names.add(name.casefold())
            worksheet = workbook.create_sheet(name)
            title = _clean_text(spec.get("title"), 500)
            row_index = 1
            if title:
                worksheet.cell(row=row_index, column=1, value=title)
                worksheet.cell(row=row_index, column=1).font = Font(size=16, bold=True, color="1F4E79")
                row_index += 2
            headers = list(spec.get("headers") or [])[:100]
            rows = list(spec.get("rows") or [])[:10_000]
            if headers:
                for column, value in enumerate(headers, 1):
                    cell = worksheet.cell(row=row_index, column=column, value=_clean_text(value, 1_000))
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E79")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                row_index += 1
            for values in rows:
                if not isinstance(values, list):
                    values = [values]
                for column, value in enumerate(values[:100], 1):
                    worksheet.cell(row=row_index, column=column, value=value)
                row_index += 1
            worksheet.freeze_panes = f"A{3 if title and headers else 2 if headers else 1}"
            for column in range(1, min(worksheet.max_column, 100) + 1):
                observed = max(
                    (len(str(worksheet.cell(row=row, column=column).value or "")) for row in range(1, min(worksheet.max_row, 200) + 1)),
                    default=8,
                )
                worksheet.column_dimensions[get_column_letter(column)].width = min(max(observed + 2, 10), 45)
        workbook.save(destination)

    try:
        _atomic_save(target, save)
        reopened = load_workbook(target, read_only=True, data_only=False)
        names = reopened.sheetnames
        reopened.close()
        if not names:
            raise ValueError("XLSX validation found no worksheets")
    except Exception as exc:
        target.unlink(missing_ok=True)
        return ToolResult(f"ERROR creating XLSX: {exc}", ok=False)
    return ToolResult(
        f"OK: Created and structurally validated XLSX ({len(names)} sheet(s): {', '.join(names)}): {target}",
        verification_ids=("artifact:xlsx",),
    )


def create_pptx(
    path: str,
    title: str,
    subtitle: str = "",
    slides: list[dict] | None = None,
    sources: list[str] | None = None,
    overwrite: bool = False,
) -> ToolResult:
    target = _target(path, ".pptx", overwrite)
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return ToolResult("ERROR: PPTX creation requires python-pptx. Install project requirements.", ok=False)
    normalized_slides = [item for item in list(slides or [])[:50] if isinstance(item, dict)]
    normalized_sources = [_clean_text(item, 2_000) for item in (sources or [])[:50] if _clean_text(item)]

    def save(destination: Path) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        cover = presentation.slides.add_slide(presentation.slide_layouts[0])
        cover.shapes.title.text = _clean_text(title, 500) or "Untitled presentation"
        cover.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
        if len(cover.placeholders) > 1:
            cover.placeholders[1].text = _clean_text(subtitle, 1_000)
        for spec in normalized_slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = _clean_text(spec.get("title"), 500) or "Section"
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
            body = slide.placeholders[1].text_frame
            body.clear()
            bullets = spec.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            for index, bullet in enumerate(list(bullets)[:15]):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = _clean_text(bullet, 2_000)
                paragraph.level = 0
                paragraph.font.size = Pt(20)
        if normalized_sources:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Sources"
            body = slide.placeholders[1].text_frame
            body.clear()
            for index, source in enumerate(normalized_sources):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = source
                paragraph.font.size = Pt(12)
        presentation.save(destination)

    try:
        _atomic_save(target, save)
        reopened = Presentation(target)
        slide_count = len(reopened.slides)
        if slide_count < 1:
            raise ValueError("PPTX validation found no slides")
    except Exception as exc:
        target.unlink(missing_ok=True)
        return ToolResult(f"ERROR creating PPTX: {exc}", ok=False)
    return ToolResult(
        f"OK: Created and structurally validated PPTX ({slide_count} slides): {target}",
        verification_ids=("artifact:pptx",),
    )


def create_pdf(
    path: str,
    title: str,
    summary: str = "",
    sections: list[dict] | None = None,
    sources: list[str] | None = None,
    overwrite: bool = False,
) -> ToolResult:
    target = _target(path, ".pdf", overwrite)
    try:
        from pypdf import PdfReader
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        return ToolResult("ERROR: PDF creation requires reportlab and pypdf. Install project requirements.", ok=False)
    normalized_sections = [item for item in list(sections or [])[:50] if isinstance(item, dict)]
    normalized_sources = [_clean_text(item, 2_000) for item in (sources or [])[:50] if _clean_text(item)]

    def save(destination: Path) -> None:
        styles = getSampleStyleSheet()
        story = [Paragraph(_clean_text(title, 500) or "Untitled report", styles["Title"]), Spacer(1, 0.18 * inch)]
        if summary_text := _clean_text(summary):
            story.extend([Paragraph(summary_text, styles["BodyText"]), Spacer(1, 0.14 * inch)])
        for item in normalized_sections:
            if heading := _clean_text(item.get("heading"), 500):
                story.append(Paragraph(heading, styles["Heading1"]))
            paragraphs = item.get("paragraphs") or []
            if isinstance(paragraphs, str):
                paragraphs = [paragraphs]
            for text in list(paragraphs)[:50]:
                if cleaned := _clean_text(text):
                    story.extend([Paragraph(cleaned, styles["BodyText"]), Spacer(1, 0.08 * inch)])
            bullets = item.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            items = [ListItem(Paragraph(_clean_text(text, 5_000), styles["BodyText"])) for text in list(bullets)[:50] if _clean_text(text)]
            if items:
                story.append(ListFlowable(items, bulletType="bullet"))
        if normalized_sources:
            story.append(Paragraph("Sources", styles["Heading1"]))
            story.append(ListFlowable([ListItem(Paragraph(source, styles["BodyText"])) for source in normalized_sources], bulletType="bullet"))
        document = SimpleDocTemplate(str(destination), pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch)
        document.build(story)

    try:
        _atomic_save(target, save)
        page_count = len(PdfReader(str(target)).pages)
        if page_count < 1:
            raise ValueError("PDF validation found no pages")
    except Exception as exc:
        target.unlink(missing_ok=True)
        return ToolResult(f"ERROR creating PDF: {exc}", ok=False)
    return ToolResult(
        f"OK: Created and structurally validated PDF ({page_count} pages): {target}",
        verification_ids=("artifact:pdf",),
    )


_SECTION_SCHEMA = {
    "type": "array",
    "items": {
        "type": "object",
        "properties": {
            "heading": {"type": "string"},
            "paragraphs": {"type": "array", "items": {"type": "string"}},
            "bullets": {"type": "array", "items": {"type": "string"}},
        },
    },
}


TOOLS = [
    CodeTool("create_docx", "Create a styled Word DOCX at an absolute path from a title, summary, structured sections, and source URLs; reopen it for structural validation.", {"type": "object", "properties": {"path": {"type": "string"}, "title": {"type": "string"}, "summary": {"type": "string"}, "sections": _SECTION_SCHEMA, "sources": {"type": "array", "items": {"type": "string"}}, "overwrite": {"type": "boolean", "default": False}}, "required": ["path", "title"]}, create_docx, "write", "artifact", ("CODE_EDIT", "RESEARCH", "FILE_OPS")),
    CodeTool("create_xlsx", "Create a styled Excel XLSX at an absolute path from worksheet names, headers, rows, values, and formulas; reopen it for structural validation.", {"type": "object", "properties": {"path": {"type": "string"}, "sheets": {"type": "array", "items": {"type": "object", "properties": {"name": {"type": "string"}, "title": {"type": "string"}, "headers": {"type": "array", "items": {"type": "string"}}, "rows": {"type": "array", "items": {"type": "array"}}}, "required": ["name", "rows"]}}, "overwrite": {"type": "boolean", "default": False}}, "required": ["path", "sheets"]}, create_xlsx, "write", "artifact", ("CODE_EDIT", "RESEARCH", "FILE_OPS")),
    CodeTool("create_pptx", "Create a widescreen PowerPoint PPTX at an absolute path from a title, subtitle, slides, bullet content, and sources; reopen it for structural validation.", {"type": "object", "properties": {"path": {"type": "string"}, "title": {"type": "string"}, "subtitle": {"type": "string"}, "slides": {"type": "array", "items": {"type": "object", "properties": {"title": {"type": "string"}, "bullets": {"type": "array", "items": {"type": "string"}}}, "required": ["title", "bullets"]}}, "sources": {"type": "array", "items": {"type": "string"}}, "overwrite": {"type": "boolean", "default": False}}, "required": ["path", "title", "slides"]}, create_pptx, "write", "artifact", ("CODE_EDIT", "RESEARCH", "FILE_OPS")),
    CodeTool("create_pdf", "Create a styled PDF report at an absolute path from a title, summary, structured sections, and source URLs; reopen it for page validation.", {"type": "object", "properties": {"path": {"type": "string"}, "title": {"type": "string"}, "summary": {"type": "string"}, "sections": _SECTION_SCHEMA, "sources": {"type": "array", "items": {"type": "string"}}, "overwrite": {"type": "boolean", "default": False}}, "required": ["path", "title"]}, create_pdf, "write", "artifact", ("CODE_EDIT", "RESEARCH", "FILE_OPS")),
]
